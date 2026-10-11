"""Gera o Word a partir do conteúdo revisável; requer python-docx e Pillow.
Execute na raiz: python docs/fontes/gerar_documentacao.py
O PDF deve ser exportado e conferido visualmente após qualquer alteração.
"""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

root=Path(__file__).resolve().parents[2]
source=Path(__file__).with_name('documentacao.md')
output=root/'docs/final/Documentacao_OBotBarato.docx'
output.parent.mkdir(parents=True,exist_ok=True)
doc=Document()
section=doc.sections[0]
section.page_width=Cm(21);section.page_height=Cm(29.7)
section.top_margin=Cm(3);section.left_margin=Cm(3)
section.bottom_margin=Cm(2);section.right_margin=Cm(2)
section.header_distance=Cm(1.3);section.footer_distance=Cm(1)
section.different_first_page_header_footer=True
normal=doc.styles['Normal']; normal.font.name='Times New Roman';normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(8)
normal.paragraph_format.first_line_indent=Cm(1.25)
normal.paragraph_format.widow_control=True
for name,size in [('Title',16),('Heading 1',14),('Heading 2',12)]:
    s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.space_before=Pt(0);s.paragraph_format.space_after=Pt(12);s.paragraph_format.keep_with_next=True
for stylename in ['Normal','Title','Title Char','Heading 1','Heading 1 Char','Heading 2','Heading 2 Char']:
    if stylename not in doc.styles:continue
    st=doc.styles[stylename]
    fonts=st._element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if 'theme' in key.lower():del fonts.attrib[key]
    for script in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+script),'Times New Roman')
    for border in st._element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
doc.core_properties.title='OBotBarato Sistema desktop de perfumaria'
doc.core_properties.author='Giovana Paula Felix dos Santos; Samuel da Silva Guimarães'
doc.core_properties.subject='Projeto acadêmico de Sistemas de Informação'
doc.core_properties.keywords='desktop, Java Swing, PostgreSQL, perfumaria'
doc.core_properties.comments=''
f=section.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.paragraph_format.first_line_indent=Cm(0)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
pending_page=False
figure_number=0
def p(text='',style=None,center=False):
    global pending_page
    x=doc.add_paragraph(text,style)
    if pending_page:x.paragraph_format.page_break_before=True;pending_page=False
    if center:x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Cm(0)
    elif style is None:x.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return x
def space(n):
    x=p('',center=True);x.paragraph_format.space_after=Pt(n);x.paragraph_format.line_spacing=1
def newpage():
    global pending_page
    pending_page=True

p('UNIVERSIDADE PARANAENSE',center=True).runs[0].bold=True
p('CURSO DE SISTEMAS DE INFORMAÇÃO',center=True)
space(48)
p('GIOVANA PAULA FELIX DOS SANTOS',center=True)
p('SAMUEL DA SILVA GUIMARÃES',center=True)
space(100)
p('OBotBarato','Title',True)
p('Sistema desktop de perfumaria',center=True)
space(140)
p('PARANAVAÍ',center=True);p('2026',center=True)
newpage()
p('GIOVANA PAULA FELIX DOS SANTOS',center=True)
p('SAMUEL DA SILVA GUIMARÃES',center=True)
space(82)
p('OBotBarato','Title',True);p('Sistema desktop de perfumaria',center=True)
space(48)
x=p('Projeto acadêmico apresentado ao Curso de Sistemas de Informação da Universidade Paranaense, em Paranavaí. Apresentação prevista para 28 de outubro de 2026.');x.paragraph_format.left_indent=Cm(7);x.paragraph_format.first_line_indent=Cm(0);x.paragraph_format.line_spacing=1
x=p('Orientação: Prof. Me. Jaime William Dias.');x.paragraph_format.left_indent=Cm(7);x.paragraph_format.first_line_indent=Cm(0);x.paragraph_format.line_spacing=1
space(160)
p('PARANAVAÍ',center=True);p('2026',center=True)
newpage()
p('Sumário','Heading 1')
content=source.read_text('utf-8')
tocpath=root/'build/finalizacao/toc.json'
toc=json.loads(tocpath.read_text('utf-8')) if tocpath.exists() else {}
chapters=[line[2:] for line in content.splitlines() if line.startswith('# ') and line!='# Resumo']
for title in chapters:
    x=p(title+'\t'+str(toc.get(title,'')));x.paragraph_format.first_line_indent=Cm(0);x.alignment=WD_ALIGN_PARAGRAPH.LEFT
    x.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    x.paragraph_format.line_spacing=1.5
newpage()

for index,page in enumerate(content.split('<!-- página -->')):
    if index:newpage()
    lines=page.strip().splitlines();i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('# '):p(line[2:],'Heading 1');i+=1
        elif line.startswith('## '):p(line[3:],'Heading 2');i+=1
        elif line.startswith('!['):
            match=re.match(r'!\[(.*?)\]\((.*?)\)',line);img=(source.parent/match[2]).resolve()
            if not img.exists():raise FileNotFoundError(img)
            with Image.open(img) as im:w,h=im.size
            max_height=15 if 'validacao-ui' in str(img) else 16.5
            width=min(16,max_height*w/h)
            x=p('',center=True);x.paragraph_format.space_after=Pt(6);x.paragraph_format.line_spacing=1
            x.add_run().add_picture(str(img),width=Cm(width))
            x._p.xpath('.//wp:docPr')[0].set('descr',match[1])
            figure_number+=1
            caption=p('Figura '+str(figure_number)+' '+match[1],center=True);caption.paragraph_format.line_spacing=1;caption.paragraph_format.space_after=Pt(6);caption.runs[0].font.size=Pt(10)
            i+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                row=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-+:?$',x) for x in row):rows.append(row)
                i+=1
            table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Table Grid';table.autofit=False
            widths=[2,14] if len(rows[0])==2 else [5.5]+[3.5]*(len(rows[0])-1)
            for column,width in zip(table.columns,widths):column.width=Cm(width)
            for rr,row in enumerate(rows):
                cells=table.add_row().cells
                trPr=table.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
                if rr==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
                for j,(cell,value) in enumerate(zip(cells,row)):
                    cell.width=Cm(widths[j]);cell.text=value
                    props=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'50' if edge in ['top','bottom'] else '90');e.set(qn('w:type'),'dxa');margin.append(e)
                    props.append(margin)
                    for x in cell.paragraphs:
                        x.paragraph_format.first_line_indent=Cm(0);x.paragraph_format.line_spacing=1;x.paragraph_format.space_after=Pt(1);x.paragraph_format.space_before=Pt(1)
                        for run in x.runs:run.font.size=Pt(10);run.bold=rr==0
                    if rr==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EEEEEE');props.append(shade)
            p('').paragraph_format.space_after=Pt(2)
        else:
            buf=[line];i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(('#','!','|')):buf.append(lines[i].strip());i+=1
            x=p(' '.join(buf))
            if buf[0].startswith('Fonte:'):x.paragraph_format.first_line_indent=Cm(0);x.paragraph_format.line_spacing=1; x.runs[0].font.size=Pt(10)
doc.save(output)
print(str(output))
